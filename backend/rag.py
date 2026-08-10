from __future__ import annotations

import psycopg2
import os
import io
import numpy as np
import asyncio
import logging
from typing import Optional
from dotenv import load_dotenv
from openai import AsyncOpenAI

logger = logging.getLogger(__name__)

load_dotenv(os.path.join(os.path.dirname(__file__), '.env'))

OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "")
EMBED_MODEL    = os.getenv("EMBED_MODEL", "text-embedding-3-small")
LLM_MODEL      = os.getenv("LLM_MODEL", "gpt-4o-mini")
CHUNK_SIZE     = 1500
CHUNK_OVERLAP  = 150
EMBEDDING_DIM  = 1536

def _get_openai_client() -> AsyncOpenAI:
    key = os.getenv("OPENAI_API_KEY", OPENAI_API_KEY)
    return AsyncOpenAI(api_key=key)

openai_client = _get_openai_client()

DB_CONFIG = {
    "host":     os.getenv("PG_HOST",     "127.0.0.1"),
    "port":     int(os.getenv("PG_PORT", 5433)),
    "user":     os.getenv("PG_USER",     "postgres"),
    "password": os.getenv("PG_PASSWORD", "mysecurepassword123"),
    "database": os.getenv("PG_DB",       "offline_db"),
    "sslmode":  os.getenv("PG_SSLMODE",  "prefer"),
}


def get_conn():
    conn = psycopg2.connect(**DB_CONFIG)
    conn.autocommit = False
    return conn


def init_vector_table():
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("CREATE EXTENSION IF NOT EXISTS vector;")
            cur.execute(f"""
                CREATE TABLE IF NOT EXISTS document_chunks (
                    id          serial PRIMARY KEY,
                    user_id     integer DEFAULT 0,
                    filename    text,
                    chunk_index int,
                    content     text,
                    embedding   vector({EMBEDDING_DIM}),
                    source_type text DEFAULT 'text'
                );
            """)
            cur.execute("ALTER TABLE document_chunks ADD COLUMN IF NOT EXISTS user_id integer DEFAULT 0;")
            cur.execute("ALTER TABLE document_chunks ADD COLUMN IF NOT EXISTS source_type text DEFAULT 'text';")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_chunks_user ON document_chunks (user_id);")
        conn.commit()


def init_upload_jobs_table():
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS upload_jobs (
                    id SERIAL PRIMARY KEY,
                    user_id INTEGER NOT NULL,
                    filename TEXT NOT NULL,
                    file_type TEXT,
                    status TEXT NOT NULL DEFAULT 'processing',
                    chunks INTEGER DEFAULT 0,
                    error TEXT,
                    created_at TIMESTAMP DEFAULT NOW(),
                    updated_at TIMESTAMP DEFAULT NOW()
                )
            """)
        conn.commit()


def create_upload_job(user_id: int, filename: str, file_type: str) -> int:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO upload_jobs (user_id, filename, file_type, status) "
                "VALUES (%s, %s, %s, 'processing') RETURNING id",
                (user_id, filename, file_type)
            )
            job_id = cur.fetchone()[0]
        conn.commit()
    return job_id


def update_upload_job(job_id: int, status: str, chunks: int = 0, error: str = None):
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "UPDATE upload_jobs SET status = %s, chunks = %s, error = %s, updated_at = NOW() WHERE id = %s",
                (status, chunks, error, job_id)
            )
        conn.commit()


def get_recent_jobs(user_id: int, limit: int = 30) -> list[dict]:
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, filename, file_type, status, chunks, error, created_at "
                "FROM upload_jobs WHERE user_id = %s ORDER BY created_at DESC LIMIT %s",
                (user_id, limit)
            )
            rows = cur.fetchall()
    return [
        {"id": r[0], "filename": r[1], "file_type": r[2], "status": r[3],
         "chunks": r[4], "error": r[5], "created_at": r[6].isoformat() if r[6] else None}
        for r in rows
    ]


def extract_text(filename: str, content: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        try:
            import fitz
            doc = fitz.open(stream=io.BytesIO(content), filetype="pdf")
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return "".join(c for c in text if c.isprintable() or c in ['\n', '\t', ' '])
        except ImportError:
            pass
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return "".join(c for c in text if c.isprintable() or c in ['\n', '\t', ' '])
        except ImportError:
            pass
    if ext in ("docx", "doc"):
        try:
            import docx
            doc = docx.Document(io.BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            pass
    if ext == "xlsx":
        max_rows  = int(os.getenv("XLSX_MAX_ROWS_PER_SHEET", "50"))
        max_cells = int(os.getenv("XLSX_MAX_TOTAL_CELLS",    "5000"))
        import openpyxl
        wb    = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        parts = []
        total = 0
        for sheet in wb.worksheets:
            parts.append(f"\n=== XLSX sheet: {getattr(sheet, 'title', 'Sheet')} ===")
            for idx, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if idx > max_rows:
                    break
                cells = ["" if c is None else str(c) for c in row]
                if not any(c.strip() for c in cells):
                    continue
                total += len(cells)
                if total > max_cells:
                    parts.append("\n[... XLSX truncated ...]")
                    return "\n".join(parts).strip()
                parts.append("\t".join(cells).rstrip("\t"))
        return "\n".join(parts).strip()
    return content.decode("utf-8", errors="ignore")


def split_text(text: str) -> list[str]:
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start:start + CHUNK_SIZE].strip())
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return [c for c in chunks if c]


def cosine_similarity(a: list[float], b: list[float]) -> float:
    va, vb = np.array(a), np.array(b)
    return float(np.dot(va, vb) / (np.linalg.norm(va) * np.linalg.norm(vb) + 1e-10))


def get_all_files_info(user_id: int = 0) -> list[dict]:
    with get_conn() as conn:
        with conn.cursor() as cur:
            if user_id > 0:
                cur.execute("""
                    SELECT filename, COUNT(*) as chunk_count, source_type
                    FROM document_chunks
                    WHERE user_id = %s AND source_type != 'pending'
                    GROUP BY filename, source_type ORDER BY filename
                """, (user_id,))
            else:
                cur.execute("""
                    SELECT filename, COUNT(*) as chunk_count, source_type
                    FROM document_chunks
                    WHERE source_type != 'pending'
                    GROUP BY filename, source_type ORDER BY filename
                """)
            rows = cur.fetchall()
    return [{"filename": r[0], "chunk_count": r[1], "source_type": r[2]} for r in rows]


async def check_ollama_health() -> tuple[bool, str]:
    try:
        await openai_client.models.list()
        return True, "ok"
    except Exception as e:
        return False, str(e)


async def get_embedding(text: str) -> list[float]:
    client = _get_openai_client()
    res = await client.embeddings.create(model=EMBED_MODEL, input=text)
    emb = res.data[0].embedding
    if not emb:
        raise ValueError("Empty embedding returned")
    return emb


def rerank_chunks(chunks: list[dict]) -> list[dict]:
    if not chunks:
        return chunks
    counts: dict[str, int] = {}
    for c in chunks:
        counts[c["filename"]] = counts.get(c["filename"], 0) + 1
    top = max(counts, key=counts.get)
    out = []
    for c in chunks:
        bonus = 0.05 if c["filename"] == top else 0.0
        out.append({**c, "rerank_score": round(c["similarity"] + bonus, 4)})
    out.sort(key=lambda x: x["rerank_score"], reverse=True)
    return out


def diverse_retrieve(all_chunks: list[dict], top_k: int) -> list[dict]:
    best_per_file: dict[str, dict] = {}
    for c in all_chunks:
        fn = c["filename"]
        if fn not in best_per_file or c["similarity"] > best_per_file[fn]["similarity"]:
            best_per_file[fn] = c
    diverse = sorted(best_per_file.values(), key=lambda x: x["similarity"], reverse=True)
    if len(diverse) < top_k:
        used = {(c["filename"], c["chunk_index"]) for c in diverse}
        for c in all_chunks:
            key = (c["filename"], c["chunk_index"])
            if key not in used:
                diverse.append(c)
                used.add(key)
            if len(diverse) >= top_k:
                break
    return diverse[:top_k]


def vector_store_info(user_id: int = 0) -> dict:
    with get_conn() as conn:
        with conn.cursor() as cur:
            if user_id > 0:
                cur.execute("SELECT COUNT(*) FROM document_chunks WHERE source_type != 'pending' AND user_id = %s", (user_id,))
                total_chunks = cur.fetchone()[0]
                cur.execute("SELECT COUNT(DISTINCT filename) FROM document_chunks WHERE source_type != 'pending' AND filename IS NOT NULL AND user_id = %s", (user_id,))
                total_docs = cur.fetchone()[0]
                cur.execute("SELECT source_type, COUNT(*) FROM document_chunks WHERE source_type != 'pending' AND user_id = %s GROUP BY source_type", (user_id,))
            else:
                cur.execute("SELECT COUNT(*) FROM document_chunks WHERE source_type != 'pending'")
                total_chunks = cur.fetchone()[0]
                cur.execute("SELECT COUNT(DISTINCT filename) FROM document_chunks WHERE source_type != 'pending' AND filename IS NOT NULL")
                total_docs = cur.fetchone()[0]
                cur.execute("SELECT source_type, COUNT(*) FROM document_chunks WHERE source_type != 'pending' GROUP BY source_type")
            by_type = {r[0]: r[1] for r in cur.fetchall()}
    return {
        "total_chunks":    total_chunks,
        "total_documents": total_docs,
        "embed_model":     EMBED_MODEL,
        "embedding_dim":   EMBEDDING_DIM,
        "by_source_type":  by_type,
    }


class RAGPipeline:

    def prepare(self, filename: str, content: bytes) -> list[str]:
        return split_text(extract_text(filename, content))

    async def embed(self, chunks: list[str]) -> list[list[float]]:
        client = _get_openai_client()
        BATCH = 512
        results = []
        for i in range(0, len(chunks), BATCH):
            batch = chunks[i:i + BATCH]
            res = await client.embeddings.create(model=EMBED_MODEL, input=batch)
            res.data.sort(key=lambda x: x.index)
            results.extend([d.embedding for d in res.data])
        return results

    async def store(self, filename: str, chunks: list[str],
                    embeddings: list[list[float]],
                    user_id: int = 0, source_type: str = "text") -> int:
        init_vector_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM document_chunks WHERE filename = %s AND user_id = %s", (filename, user_id))
                for i, (chunk, emb) in enumerate(zip(chunks, embeddings)):
                    cur.execute(
                        "INSERT INTO document_chunks (user_id, filename, chunk_index, content, embedding, source_type) "
                        "VALUES (%s, %s, %s, %s, %s, %s)",
                        (user_id, filename, i, chunk, emb, source_type)
                    )
            conn.commit()
        return len(chunks)

    async def embed_all(self, filename: str, user_id: int = 0) -> None:
        from psycopg2.extras import execute_batch
        _ensure_indexing_table()
        try:
            with get_conn() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        "SELECT id, content FROM document_chunks "
                        "WHERE filename = %s AND user_id = %s AND embedding IS NULL ORDER BY chunk_index",
                        (filename, user_id)
                    )
                    rows = cur.fetchall()
                    cur.execute(
                        "SELECT COUNT(*) FROM document_chunks WHERE filename = %s AND user_id = %s",
                        (filename, user_id)
                    )
                    total = cur.fetchone()[0]
            if not rows:
                clear_indexing_progress(user_id, filename)
                return
            ids   = [r[0] for r in rows]
            texts = [r[1] for r in rows]
            client = _get_openai_client()
            BATCH = 512
            batches = [texts[i:i+BATCH] for i in range(0, len(texts), BATCH)]
            results = await asyncio.gather(*[
                client.embeddings.create(model=EMBED_MODEL, input=b) for b in batches
            ])
            embeddings: list[list[float]] = []
            for res in results:
                res.data.sort(key=lambda x: x.index)
                embeddings.extend([d.embedding for d in res.data])
            with get_conn() as conn:
                with conn.cursor() as cur:
                    execute_batch(
                        cur,
                        "UPDATE document_chunks SET embedding = %s::vector WHERE id = %s",
                        [(embeddings[i], ids[i]) for i in range(len(ids))],
                        page_size=200
                    )
                conn.commit()
            set_indexing_progress(user_id, filename, total, total, "done")
            clear_indexing_progress(user_id, filename)
        except Exception as e:
            logger.error(f"embed_all error for {filename}: {e}")
            set_indexing_progress(user_id, filename, 0, 0, "error")

    async def embed_one_batch(self, filename: str, user_id: int = 0, batch_size: int = 50) -> dict:
        """Embed one batch of un-embedded chunks. Called repeatedly by frontend."""
        from psycopg2.extras import execute_values
        _ensure_indexing_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id, content FROM document_chunks "
                    "WHERE filename = %s AND user_id = %s AND embedding IS NULL "
                    "ORDER BY chunk_index LIMIT %s",
                    (filename, user_id, batch_size)
                )
                rows = cur.fetchall()
                cur.execute(
                    "SELECT COUNT(*) FROM document_chunks WHERE filename = %s AND user_id = %s",
                    (filename, user_id)
                )
                total = cur.fetchone()[0]
                cur.execute(
                    "SELECT COUNT(*) FROM document_chunks WHERE filename = %s AND user_id = %s AND embedding IS NOT NULL",
                    (filename, user_id)
                )
                done_before = cur.fetchone()[0]

        if not rows:
            clear_indexing_progress(user_id, filename)
            return {"filename": filename, "total": total, "done": total, "status": "done"}

        ids, texts = zip(*rows)
        embeddings = await self.embed(list(texts))

        with get_conn() as conn:
            with conn.cursor() as cur:
                execute_values(
                    cur,
                    "UPDATE document_chunks SET embedding = data.emb::vector "
                    "FROM (VALUES %s) AS data(id, emb) WHERE document_chunks.id = data.id",
                    [(ids[i], str(embeddings[i])) for i in range(len(ids))]
                )
            conn.commit()

        done_now = done_before + len(ids)
        remaining = total - done_now
        status = "done" if remaining <= 0 else "indexing"
        set_indexing_progress(user_id, filename, total, done_now, status)
        if status == "done":
            clear_indexing_progress(user_id, filename)
        return {"filename": filename, "total": total, "done": done_now, "status": status}

    async def store_fast(self, filename: str, chunks: list[str],
                         user_id: int = 0, source_type: str = "text") -> int:
        from psycopg2.extras import execute_values
        init_vector_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("DELETE FROM document_chunks WHERE filename = %s AND user_id = %s", (filename, user_id))
                rows = [
                    (user_id, filename, i,
                     "".join(c for c in chunk if c.isprintable() or c in ['\n', '\t', ' ']),
                     None, source_type)
                    for i, chunk in enumerate(chunks)
                ]
                execute_values(cur,
                    "INSERT INTO document_chunks (user_id, filename, chunk_index, content, embedding, source_type) VALUES %s",
                    rows)
            conn.commit()
        return len(chunks)

    async def retrieve(self, query: str, top_k: int = 10, user_id: int = 0) -> list[dict]:
        emb     = await get_embedding(query)
        fetch_k = max(top_k * 6, 60)
        with get_conn() as conn:
            with conn.cursor() as cur:
                if user_id > 0:
                    cur.execute("""
                        SELECT filename, chunk_index, content, source_type,
                               1 - (embedding <=> %s::vector) AS similarity
                        FROM document_chunks
                        WHERE user_id = %s AND embedding IS NOT NULL
                        ORDER BY embedding <=> %s::vector LIMIT %s
                    """, (emb, user_id, emb, fetch_k))
                else:
                    cur.execute("""
                        SELECT filename, chunk_index, content, source_type,
                               1 - (embedding <=> %s::vector) AS similarity
                        FROM document_chunks
                        WHERE embedding IS NOT NULL
                        ORDER BY embedding <=> %s::vector LIMIT %s
                    """, (emb, emb, fetch_k))
                rows = cur.fetchall()
        if not rows:
            return []
        all_chunks = [
            {"filename": r[0], "chunk_index": r[1], "content": r[2],
             "source_type": r[3], "similarity": round(r[4], 4)}
            for r in rows
        ]
        return rerank_chunks(diverse_retrieve(all_chunks, top_k))

    def augment_prompt(self, query: str, chunks: list[dict],
                       all_files: Optional[list[dict]] = None,
                       total_docs: int = 0, total_chunks: int = 0) -> str:
        if all_files:
            files_list = "\n".join(
                f"  {i+1}. {f['filename']} ({f['chunk_count']} чанков)"
                for i, f in enumerate(all_files)
            )
            system_context = f"ФАЙЛЫ В БАЗЕ ({total_docs} шт.):\n{files_list}\n"
        else:
            system_context = f"Файлов в базе: {total_docs} | Чанков: {total_chunks}\n"

        context = "\n\n".join(
            f"[{c['filename']} | чанк {c['chunk_index']}]\n{c['content']}"
            for c in chunks
        )
        return (
            f"Ты — умный универсальный ИИ-ассистент базы знаний. Отвечай на русском языке.\n\n"
            f"ПРАВИЛА:\n"
            f"1. Если в документах есть ответ — используй их как основу, дополняй своими знаниями.\n"
            f"2. Если в документах нет ответа — отвечай из своих знаний, не говори что не можешь помочь.\n"
            f"3. На вопросы про количество/список файлов — отвечай по ПОЛНОМУ СПИСКУ выше.\n"
            f"4. Используй Markdown: **жирный**, ## заголовки, - списки.\n"
            f"5. Никогда не пиши 'представлен выше' — давай конкретный ответ.\n\n"
            f"{'='*50}\n"
            f"{system_context}"
            f"{'='*50}\n\n"
            f"ФРАГМЕНТЫ ПО ЗАПРОСУ:\n{context}\n\n"
            f"ВОПРОС: {query}\n\nОТВЕТ:"
        )

    async def generate(self, prompt: str) -> str:
        client = _get_openai_client()
        res = await client.chat.completions.create(
            model=LLM_MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
            max_tokens=2048,
        )
        return res.choices[0].message.content.strip()

    async def evaluate(self, query: str, answer: str) -> float:
        try:
            q_emb = await get_embedding(query)
            a_emb = await get_embedding(answer)
            return round(cosine_similarity(q_emb, a_emb), 4)
        except Exception:
            return 0.0

    async def run(self, query: str, top_k: int = 10, user_id: int = 0) -> dict:
        greetings = ['привет', 'hello', 'hi', 'здравствуй', 'здравствуйте',
                     'добрый день', 'добрый вечер', 'доброе утро', 'хай', 'салют']
        if query.lower().strip().rstrip('!.,') in greetings:
            return {
                "query": query,
                "answer": "Привет! Я готов помочь. Задайте любой вопрос — отвечу по документам из базы знаний или из своих знаний.",
                "sources": [], "cosine_similarity": 1.0,
            }

        info      = vector_store_info(user_id)
        all_files = get_all_files_info(user_id)
        chunks    = await self.retrieve(query, top_k, user_id)

        if not chunks:
            prompt = (
                f"Ты — умный универсальный ИИ-ассистент. Отвечай развёрнуто, полезно и точно на русском языке.\n"
                f"База знаний пуста или не содержит релевантных документов по этому вопросу.\n"
                f"Отвечай из своих знаний — не говори что не можешь помочь.\n\n"
                f"Вопрос: {query}\n\nОтвет:"
            )
            answer = await self.generate(prompt)
            return {"query": query, "answer": answer, "sources": [], "cosine_similarity": 0.0}

        prompt = self.augment_prompt(query, chunks, all_files=all_files,
                                     total_docs=info["total_documents"],
                                     total_chunks=info["total_chunks"])
        answer = await self.generate(prompt)
        score  = await self.evaluate(query, answer)

        return {
            "query":  query,
            "answer": answer,
            "sources": [
                {"filename": c["filename"], "chunk_index": c["chunk_index"],
                 "similarity": c["similarity"], "rerank_score": c.get("rerank_score")}
                for c in chunks
            ],
            "source_chunks": [
                {"filename": c["filename"], "content": c["content"][:500], "similarity": c["similarity"]}
                for c in chunks
            ],
            "cosine_similarity": score,
        }

        prompt = self.augment_prompt(query, chunks, all_files=all_files,
                                     total_docs=info["total_documents"],
                                     total_chunks=info["total_chunks"])
        answer = await self.generate(prompt)
        score  = await self.evaluate(query, answer)

        return {
            "query":  query,
            "answer": answer,
            "sources": [
                {"filename": c["filename"], "chunk_index": c["chunk_index"],
                 "similarity": c["similarity"], "rerank_score": c.get("rerank_score")}
                for c in chunks
            ],
            "source_chunks": [
                {"filename": c["filename"], "content": c["content"][:500], "similarity": c["similarity"]}
                for c in chunks
            ],
            "cosine_similarity": score,
        }


pipeline = RAGPipeline()


def _ensure_indexing_table():
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS indexing_jobs (
                    user_id      integer NOT NULL DEFAULT 0,
                    filename     text NOT NULL,
                    total_chunks integer NOT NULL DEFAULT 0,
                    done_chunks  integer NOT NULL DEFAULT 0,
                    status       text NOT NULL DEFAULT 'indexing',
                    updated_at   timestamp DEFAULT now(),
                    PRIMARY KEY (user_id, filename)
                );
            """)
        conn.commit()


def get_indexing_progress(user_id: int = 0) -> dict:
    try:
        _ensure_indexing_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT filename, total_chunks, done_chunks, status "
                    "FROM indexing_jobs WHERE user_id = %s AND status IN ('indexing','error')",
                    (user_id,)
                )
                rows = cur.fetchall()
        return {r[0]: {"total": r[1], "done": r[2], "status": r[3]} for r in rows}
    except Exception:
        return {}


def set_indexing_progress(user_id: int, filename: str, total: int, done: int, status: str = "indexing"):
    try:
        _ensure_indexing_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute("""
                    INSERT INTO indexing_jobs (user_id, filename, total_chunks, done_chunks, status, updated_at)
                    VALUES (%s, %s, %s, %s, %s, now())
                    ON CONFLICT (user_id, filename) DO UPDATE SET
                        total_chunks = EXCLUDED.total_chunks,
                        done_chunks  = EXCLUDED.done_chunks,
                        status       = EXCLUDED.status,
                        updated_at   = now()
                """, (user_id, filename, total, done, status))
            conn.commit()
    except Exception:
        pass


def clear_indexing_progress(user_id: int, filename: str):
    try:
        _ensure_indexing_table()
        with get_conn() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "UPDATE indexing_jobs SET status = 'done' WHERE user_id = %s AND filename = %s",
                    (user_id, filename)
                )
            conn.commit()
    except Exception:
        pass


async def ingest_document(filename: str, content: bytes, user_id: int = 0) -> int:
    ext         = filename.rsplit(".", 1)[-1].lower()
    source_type = ext if ext in ("pdf", "docx", "txt") else "text"
    chunks      = pipeline.prepare(filename, content)
    embeddings  = await pipeline.embed(chunks)
    return await pipeline.store(filename, chunks, embeddings, user_id, source_type)


async def search_documents(query: str, top_k: int = 10, user_id: int = 0) -> list[dict]:
    return await pipeline.retrieve(query, top_k, user_id)


async def generate_answer(query: str, top_k: int = 10, user_id: int = 0) -> dict:
    return await pipeline.run(query, top_k, user_id)


def simple_rag(query: str, top_k: int = 5, user_id: int = 0) -> list[dict]:
    if user_id <= 0:
        return []
    keywords = set(query.lower().split())
    with get_conn() as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT filename, chunk_index, content FROM document_chunks WHERE user_id = %s", (user_id,))
            rows = cur.fetchall()
    scored = [
        {"filename": fn, "chunk_index": ci, "content": ct,
         "score": len(keywords & set(ct.lower().split()))}
        for fn, ci, ct in rows
        if len(keywords & set(ct.lower().split())) > 0
    ]
    scored.sort(key=lambda x: x["score"], reverse=True)
    return scored[:top_k]


async def modular_rag(query: str, top_k: int = 5, user_id: int = 0) -> list[dict]:
    keyword_results = simple_rag(query, top_k, user_id)
    vector_results  = await search_documents(query, top_k, user_id)
    seen, merged    = set(), []
    for r in keyword_results + vector_results:
        key = (r["filename"], r["chunk_index"])
        if key not in seen:
            seen.add(key)
            merged.append(r)
    return merged[:top_k]
